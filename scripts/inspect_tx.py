import openpyxl
from collections import Counter

wb = openpyxl.load_workbook(
    "/workspace/well-lab-index/data/sources/tx-pws-lab-map-table-2026-03-26.xlsx",
    data_only=True,
)
ws = wb["Public Labs"]
print("rows", ws.max_row, "cols", ws.max_column)
rows = list(ws.iter_rows(values_only=True))
for i, r in enumerate(rows[:6], 1):
    print(i, r)

header = None
data_start = None
for i, r in enumerate(rows):
    if r and str(r[0]).strip() == "Map ID":
        header = r
        data_start = i + 1
        print("header at", i + 1, r)
        break

labs = []
for r in rows[data_start:]:
    if not r or r[0] is None:
        continue
    try:
        mid = int(r[0])
    except Exception:
        print("skip", r[:3])
        continue
    labs.append(r)

print("lab count", len(labs))
print("first", labs[0])
print("last", labs[-1])

omit_kw = [
    "Water Supply",
    "Water Authority",
    "Water District",
    "Municipal",
    "Utility",
    "City of ",
    "Town of ",
    "Village of ",
    "Public Works",
    "Water Department",
    "Water System",
    "Water Works",
    "Water Resources",
    "River Authority",
    " MUD",
    " SUD",
    " WSC",
    "WSC ",
]
cats = Counter()
for r in labs:
    n = str(r[1])
    flagged = any(k.lower() in n.lower() for k in omit_kw)
    health = "health" in n.lower()
    if health:
        tag = "HEALTH"
    elif flagged:
        tag = "UTIL"
    else:
        tag = "COMM"
    cats[tag] += 1
    print(
        f"{r[0]:3} {tag:6} Micro={r[2]} LCR={r[3]} WQP={r[4]} | {n} | {r[6]}, {r[7]} {r[8]} {r[9]}"
    )
print(cats)
